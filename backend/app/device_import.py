
from io import BytesIO
from openpyxl import load_workbook

from sqlalchemy.orm import Session
from app.model import DeviceCategory, Device


# These are the official Excel template headers.
# The Excel file must use these exact column names.
REQUIRED_COLUMNS = [
    "device_name",
    "host",
    "port",
    "connection_type",
    "username",
    "password",
    "device_category_id",
]

OPTIONAL_COLUMNS = [
    "password_rotation_days",
    "comments",
    "device_status",
    "device_condition",
]

ALL_COLUMNS = REQUIRED_COLUMNS + OPTIONAL_COLUMNS


def read_excel_file(file_bytes: bytes):
    """
    Read an uploaded Excel workbook and return its headers and rows.

    Worksheet row numbers are preserved for validation reporting.
    """

    try:
        workbook = load_workbook(
            filename=BytesIO(file_bytes),
            read_only=True,
            data_only=True,
        )
    except Exception:
        return {
            "success": False,
            "errors": [
                {
                    "row": None,
                    "field": "file",
                    "reason": "Unable to read the Excel file.",
                }
            ],
            "rows": [],
        }

    worksheet = workbook.active

    if worksheet is None:
        workbook.close()

        return {
            "success": False,
            "errors": [
                {
                    "row": None,
                    "field": "file",
                    "reason": "The workbook has no active worksheet.",
                }
            ],
            "rows": [],
        }

    header_cells = next(
        worksheet.iter_rows(
            min_row=1,
            max_row=1,
            values_only=True,
        ),
        None,
    )

    if not header_cells:
        workbook.close()

        return {
            "success": False,
            "errors": [
                {
                    "row": 1,
                    "field": "header",
                    "reason": "The Excel file is empty.",
                }
            ],
            "rows": [],
        }

    headers = [
        str(value).strip() if value is not None else ""
        for value in header_cells
    ]

    errors = []

    # Check required headers.
    for column in REQUIRED_COLUMNS:
        if column not in headers:
            errors.append({
                "row": 1,
                "field": column,
                "reason": f"Missing required column: {column}",
            })

    # Check duplicate headers.
    seen_headers = set()

    for header in headers:
        if header and header in seen_headers:
            errors.append({
                "row": 1,
                "field": header,
                "reason": f"Duplicate column header: {header}",
            })

        if header:
            seen_headers.add(header)

    # Check unsupported headers.
    for header in headers:
        if header and header not in ALL_COLUMNS:
            errors.append({
                "row": 1,
                "field": header,
                "reason": f"Unsupported column: {header}",
            })

    if errors:
        workbook.close()

        return {
            "success": False,
            "errors": errors,
            "rows": [],
        }

    rows = []

    for row_number, values in enumerate(
        worksheet.iter_rows(
            min_row=2,
            values_only=True,
        ),
        start=2,
    ):
        # Ignore completely empty rows.
        if all(value is None for value in values):
            continue

        row_data = {}

        for index, header in enumerate(headers):
            if header:
                row_data[header] = (
                    values[index]
                    if index < len(values)
                    else None
                )

        rows.append({
            "excel_row": row_number,
            "data": row_data,
        })

    workbook.close()

    return {
        "success": True,
        "errors": [],
        "rows": rows,
    }




# Allowed password rotation periods, matching schema.py
ALLOWED_ROTATION_DAYS = {90, 180, 270, 360}


def validate_device_rows(rows: list, db: Session):
    """
    Validate all parsed device rows.

    This function does not insert or update any database records.
    It returns validated rows and all validation errors.
    """

    errors = []
    validated_rows = []

    # Load existing device categories once.
    categories = {
        category.id
        for category in db.query(DeviceCategory.id).all()
    }

    for item in rows:
        row_number = item["excel_row"]
        data = item["data"]

        row_errors = []

        # Required text fields
        required_text_fields = [
            "device_name",
            "host",
            "connection_type",
            "username",
            "password",
        ]

        for field in required_text_fields:
            value = data.get(field)

            if value is None or (
                isinstance(value, str) and not value.strip()
            ):
                row_errors.append({
                    "row": row_number,
                    "field": field,
                    "reason": f"{field} is required.",
                })

            elif not isinstance(value, str):
                row_errors.append({
                    "row": row_number,
                    "field": field,
                    "reason": f"{field} must be text.",
                })

        # Device name length
        device_name = data.get("device_name")

        if isinstance(device_name, str):
            if len(device_name.strip()) > 255:
                row_errors.append({
                    "row": row_number,
                    "field": "device_name",
                    "reason": "Device name cannot exceed 255 characters.",
                })

        # Host length
        host = data.get("host")

        if isinstance(host, str):
            if len(host.strip()) > 255:
                row_errors.append({
                    "row": row_number,
                    "field": "host",
                    "reason": "Host cannot exceed 255 characters.",
                })

        # Connection type length
        connection_type = data.get("connection_type")

        if isinstance(connection_type, str):
            if len(connection_type.strip()) > 50:
                row_errors.append({
                    "row": row_number,
                    "field": "connection_type",
                    "reason": "Connection type cannot exceed 50 characters.",
                })

        # Username length
        username = data.get("username")

        if isinstance(username, str):
            if len(username.strip()) > 255:
                row_errors.append({
                    "row": row_number,
                    "field": "username",
                    "reason": "Username cannot exceed 255 characters.",
                })

        # Port validation
        port = data.get("port")

        if isinstance(port, bool):
            row_errors.append({
                "row": row_number,
                "field": "port",
                "reason": "Port must be a whole number between 1 and 65535.",
            })

        elif isinstance(port, float) and not port.is_integer():
            row_errors.append({
                "row": row_number,
                "field": "port",
                "reason": "Port must be a whole number.",
            })

        else:
            try:
                if isinstance(port, str):
                    port = int(port.strip())
                elif isinstance(port, float):
                    port = int(port)

                if not isinstance(port, int):
                    raise ValueError

                if not 1 <= port <= 65535:
                    raise ValueError

                data["port"] = port

            except (ValueError, TypeError):
                row_errors.append({
                    "row": row_number,
                    "field": "port",
                    "reason": "Port must be a whole number between 1 and 65535.",
                })

        # Device category validation
        category_id = data.get("device_category_id")

        if isinstance(category_id, bool):
            row_errors.append({
                "row": row_number,
                "field": "device_category_id",
                "reason": "Device category ID must be a whole number.",
            })

        else:
            try:
                if isinstance(category_id, str):
                    category_id = int(category_id.strip())
                elif isinstance(category_id, float):
                    if not category_id.is_integer():
                        raise ValueError
                    category_id = int(category_id)

                if not isinstance(category_id, int):
                    raise ValueError

                if category_id not in categories:
                    row_errors.append({
                        "row": row_number,
                        "field": "device_category_id",
                        "reason": "Device category does not exist.",
                    })
                else:
                    data["device_category_id"] = category_id

            except (ValueError, TypeError):
                row_errors.append({
                    "row": row_number,
                    "field": "device_category_id",
                    "reason": "Device category ID must be a valid whole number.",
                })

        # Password rotation validation
        rotation_days = data.get("password_rotation_days")

        if rotation_days is not None and rotation_days != "":
            try:
                if isinstance(rotation_days, bool):
                    raise ValueError

                if isinstance(rotation_days, str):
                    rotation_days = int(rotation_days.strip())
                elif isinstance(rotation_days, float):
                    if not rotation_days.is_integer():
                        raise ValueError
                    rotation_days = int(rotation_days)

                if rotation_days not in ALLOWED_ROTATION_DAYS:
                    raise ValueError

                data["password_rotation_days"] = rotation_days

            except (ValueError, TypeError):
                row_errors.append({
                    "row": row_number,
                    "field": "password_rotation_days",
                    "reason": "Rotation days must be 90, 180, 270, or 360.",
                })

        else:
            data["password_rotation_days"] = None

        # Optional comments
        comments = data.get("comments")

        if comments is not None and not isinstance(comments, str):
            row_errors.append({
                "row": row_number,
                "field": "comments",
                "reason": "Comments must be text.",
            })

        # Optional status and condition
        for field, default_value in [
            ("device_status", "Active"),
            ("device_condition", "Unused"),
        ]:
            value = data.get(field)

            if value is None or (
                isinstance(value, str) and not value.strip()
            ):
                data[field] = default_value

            elif not isinstance(value, str):
                row_errors.append({
                    "row": row_number,
                    "field": field,
                    "reason": f"{field} must be text.",
                })

        # Check field lengths for optional values
        for field, max_length in [
            ("device_status", 50),
            ("device_condition", 50),
        ]:
            value = data.get(field)

            if isinstance(value, str) and len(value.strip()) > max_length:
                row_errors.append({
                    "row": row_number,
                    "field": field,
                    "reason": f"{field} cannot exceed {max_length} characters.",
                })

        errors.extend(row_errors)

        if not row_errors:
            validated_rows.append(item)

    return {
        "success": len(errors) == 0,
        "errors": errors,
        "rows": validated_rows,
    }



def normalize_duplicate_value(value):
    """
    Normalize device name or host for duplicate comparison.
    Ignores leading/trailing spaces and letter case.
    """
    if value is None:
        return ""

    return str(value).strip().casefold()





def detect_device_duplicates(rows: list, db: Session):
    """
    Detect duplicate device_name + host combinations.

    Separates:
    1. Duplicate identities within the uploaded Excel file.
    2. Duplicate identities against existing database devices.

    Does not insert, update, or delete records.
    """

    within_file_duplicates = []
    database_conflicts = []

    # Track the first occurrence of each identity in Excel.
    seen_in_excel = {}

    # Collect unique identities from the workbook.
    excel_identities = set()

    for item in rows:
        row_number = item["excel_row"]
        data = item["data"]

        device_name = normalize_duplicate_value(
            data.get("device_name")
        )
        host = normalize_duplicate_value(
            data.get("host")
        )

        identity = (device_name, host)

        if identity in seen_in_excel:
            first_row = seen_in_excel[identity]

            within_file_duplicates.append({
                "type": "within_file",
                "rows": [first_row, row_number],
                "device_name": data.get("device_name"),
                "host": data.get("host"),
            })
        else:
            seen_in_excel[identity] = row_number

        excel_identities.add(identity)

    # Load existing device identities from the database.
    existing_devices = db.query(
        Device.id,
        Device.device_name,
        Device.host,
        Device.port,
        Device.connection_type,
        Device.username,
        Device.device_category_id,
        Device.password_rotation_days,
        Device.comments,
        Device.device_status,
        Device.device_condition,
    ).all()


    existing_identities = {}

    for device in existing_devices:
        identity = (
            normalize_duplicate_value(device.device_name),
            normalize_duplicate_value(device.host),
        )

        existing_identities.setdefault(
            identity, []
        ).append(device.id)

    # Find matches against existing database devices.
    for item in rows:
        row_number = item["excel_row"]
        data = item["data"]

        identity = (
            normalize_duplicate_value(
                data.get("device_name")
            ),
            normalize_duplicate_value(
                data.get("host")
            ),
        )

        matching_device_ids = existing_identities.get(
            identity, []
        )

        if matching_device_ids:
            matching_devices = [
                device
                for device in existing_devices
                if device.id in matching_device_ids
            ]

            database_conflicts.append({
                "type": "existing_database",
                "rows": [row_number],
                "device_name": data.get("device_name"),
                "host": data.get("host"),
                "existing_device_ids": matching_device_ids,
                "existing_devices": [
                    {
                        "id": device.id,
                        "device_name": device.device_name,
                        "host": device.host,
                        "port": device.port,
                        "connection_type": device.connection_type,
                        "username": device.username,
                        "device_category_id": (
                            device.device_category_id
                        ),
                        "password_rotation_days": (
                            device.password_rotation_days
                        ),
                        "comments": device.comments,
                        "device_status": device.device_status,
                        "device_condition": device.device_condition,
                    }
                    for device in matching_devices
                ],
            })



    # Keep the existing API response compatible for now.
    all_duplicates = (
        within_file_duplicates + database_conflicts
    )

    return {
        # Keep false for either kind of duplicate until
        # the commit endpoint is updated in the next step.
        "success": len(all_duplicates) == 0,
        "duplicates": all_duplicates,

        # New separate lists for the upcoming resolution flow.
        "within_file_duplicates": within_file_duplicates,
        "database_conflicts": database_conflicts,
    }

def normalize_import_device_values(device_data):
    value_maps = {
        "connection_type": {
            "ssh": "SSH",
            "rdp": "RDP",
            "ftp": "FTP",
            "sftp": "SFTP",
            "http": "HTTP",
            "https": "HTTPS",
        },
        "device_status": {
            "active": "Active",
            "inactive": "Inactive",
        },
        "device_condition": {
            "reachable": "Reachable",
            "unreachable": "Unreachable",
            "switched off": "Switched Off",
            "unused": "Unused",
        },
    }

    for field, mapping in value_maps.items():
        value = device_data.get(field)

        if isinstance(value, str):
            normalized = mapping.get(
                value.strip().casefold()
            )

            if normalized is not None:
                device_data[field] = normalized

    return device_data